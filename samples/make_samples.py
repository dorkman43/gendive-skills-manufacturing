"""gendive-skills-manufacturing 실습용 가상 데이터 생성 (난수 seed 고정).

실행: python3 samples/make_samples.py
출력: repos/gendive-skills-manufacturing/samples/
  - 생산실적.xlsx         (2026-01~09, 라인 3개 x 주·야 2교대, 품목 5개, 월·수·금만 기록)
  - 재고_출고이력.xlsx     (품목 18개, 2024-07~2026-09 일별 출고 + 품목정보·현재고·입고예정·재고0기간)
  - README.md              (함정 목록 + 검산용 정답, 이 스크립트가 계산해서 기록)
공휴일에는 생산·출고가 없다(관공서의 공휴일에 관한 규정에 따른 공휴일 표). 모든 회사·품목·숫자는 가상이다.
"""
import math
import random
import statistics
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).resolve().parent
SEED = 20261015

HOLIDAYS = {date(*d) for d in [
    (2024, 8, 15), (2024, 9, 16), (2024, 9, 17), (2024, 9, 18), (2024, 10, 1), (2024, 10, 3), (2024, 10, 9), (2024, 12, 25),
    (2025, 1, 1), (2025, 1, 27), (2025, 1, 28), (2025, 1, 29), (2025, 1, 30), (2025, 3, 3), (2025, 5, 5), (2025, 5, 6),
    (2025, 6, 3), (2025, 6, 6), (2025, 8, 15), (2025, 10, 3), (2025, 10, 6), (2025, 10, 7), (2025, 10, 8), (2025, 10, 9),
    (2025, 12, 25),
    (2026, 1, 1), (2026, 2, 16), (2026, 2, 17), (2026, 2, 18), (2026, 3, 2), (2026, 5, 1), (2026, 5, 5), (2026, 5, 25),
    (2026, 6, 3), (2026, 7, 17), (2026, 8, 17), (2026, 9, 24), (2026, 9, 25),
]}


def is_workday(d):
    return d.weekday() < 5 and d not in HOLIDAYS


def daterange(a, b):
    d = a
    while d <= b:
        yield d
        d += timedelta(days=1)


# =============================================================== 1. 생산실적
ITEMS = [  # 코드, 품명, 기준 사이클타임(초/개), 소재 구분
    ('M-101', '모터 하우징', 38, 'ST'),
    ('M-102', '커넥터 케이스', 22, 'PL'),
    ('M-103', '펌프 커버', 45, 'PL'),
    ('M-104', '기어 블록', 60, 'ST'),
    ('M-105', '고정 브라켓', 15, 'ST'),
]
LINES = ['1호라인', '2호라인', '3호라인']
SHIFTS = ['주간', '야간']
LOAD_MIN = 420  # 교대당 부하시간(근무 480분 - 식사·계획정지 60분)
TYPES = ['치수불량', '스크래치', '버(거스러미)', '기포', '도장불량', '이물', '조립불량']
TYPE_W = [0.20, 0.22, 0.15, 0.10, 0.12, 0.12, 0.09]
DOWN_REASONS = ['설비고장', '금형교체', '자재대기', '품질조정', '순간정지 누적']
BAD_LOT = 'PL-2605-B'


def lot_of(d, family):
    return f"{family}-{d.strftime('%y%m')}-{'A' if d.day <= 15 else 'B'}"


def make_production(rng):
    rows = []
    days = [d for d in daterange(date(2026, 1, 1), date(2026, 9, 30)) if is_workday(d) and d.weekday() in (0, 2, 4)]
    for d in days:
        wk = d.isocalendar()[1]
        for li, line in enumerate(LINES):
            for si, shift in enumerate(SHIFTS):
                code, name, ct, fam = ITEMS[(wk + li * 2 + si) % 5]
                down = max(0, int(rng.gauss(35, 18)))
                reason = rng.choice(DOWN_REASONS) if down >= 15 else ''
                if (d, line, shift) == (date(2026, 2, 20), '2호라인', '주간'):
                    down, reason = 240, '설비고장'
                if down < 15:
                    down = 0
                run = LOAD_MIN - down
                perf = rng.uniform(0.86, 0.95)
                prod = int(run * 60 / ct * perf)
                plan = int(LOAD_MIN * 60 / ct * 0.92 / 10) * 10
                lot = lot_of(d, fam)
                base = 0.010 if shift == '주간' else 0.014
                planted, extra = None, 0.0
                if line == '2호라인' and shift == '야간' and date(2026, 6, 1) <= d <= date(2026, 7, 15):
                    planted, extra = '치수불량', 0.055
                elif lot == BAD_LOT:
                    planted, extra = '기포', 0.05
                elif (d, line, shift) == (date(2026, 3, 18), '3호라인', '주간'):
                    planted, extra = '이물', 0.14
                rate = base + extra
                defects = max(0, int(round(prod * rate * rng.uniform(0.7, 1.3))))
                dtype = ''
                if defects:
                    if planted and rng.random() < extra / rate:
                        dtype = planted
                    else:
                        dtype = rng.choices(TYPES, TYPE_W)[0]
                good = prod - defects
                rework = int(round(prod * 0.012 * rng.uniform(0.3, 1.7)))
                rows.append(dict(d=d, line=line, shift=shift, code=code, name=name, plan=plan, prod=prod, good=good,
                                 defects=defects, dtype=dtype, rework=rework, run=run, down=down, reason=reason,
                                 lot=lot, ct=ct, fam=fam))
    # 지저분하게: 사유 빈칸, 재작업 빈칸, 일자 문자열, 라인 표기 흔들림
    cand = [r for r in rows if r['down'] > 0]
    for r in rng.sample(cand, 6):
        r['reason'] = ''
        r['reason_blank'] = True
    for r in rng.sample(rows, 3):
        r['rework'] = None
    for r in rng.sample(rows, 5):
        r['d_text'] = r['d'].strftime('%Y.%m.%d')
    for r in rng.sample(rows, 8):
        r['line_label'] = r['line'].replace('호라인', '라인')
    return rows


def write_production(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = '생산실적'
    ws.append(['일자', '라인', '교대', '품목코드', '품명', '계획수량', '생산수량', '양품수량', '불량수량', '불량유형(주된 유형)',
               '재작업수량', '가동시간(분)', '비가동시간(분)', '비가동사유', '자재로트'])
    for r in rows:
        ws.append([r.get('d_text', r['d']), r.get('line_label', r['line']), r['shift'], r['code'], r['name'], r['plan'],
                   r['prod'], r['good'], r['defects'], r['dtype'] or None, r['rework'], r['run'], r['down'],
                   r['reason'] or None, r['lot']])
        if 'd_text' not in r:
            ws.cell(row=ws.max_row, column=1).number_format = 'yyyy-mm-dd'
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws2 = wb.create_sheet('품목기준')
    ws2.append(['품목코드', '품명', '기준 사이클타임(초/개)', '소재 구분'])
    for c, n, ct, fam in ITEMS:
        ws2.append([c, n, ct, '수지(PL)' if fam == 'PL' else '철강(ST)'])
    ws3 = wb.create_sheet('안내')
    for t in ['가상 데이터입니다. 월·수·금(공휴일 제외)만 기록했습니다.',
              f'교대당 부하시간 {LOAD_MIN}분(근무 480분에서 식사·계획정지 60분 제외). 가동시간 = 부하시간 - 비가동시간.',
              '생산수량 = 양품수량 + 불량수량. 양품수량은 재작업 후 합격한 것을 포함합니다(재작업수량은 그중 재작업을 거친 수).',
              '불량유형은 그 줄에서 가장 많은 유형 하나만 적었습니다.']:
        ws3.append([t])
    wb.save(OUT / '생산실적.xlsx')


def analyze_production(rows):
    ans = {}
    tot = lambda key, rs=rows: sum(r[key] for r in rs)
    ans['n'] = len(rows)
    ans['prod'], ans['good'], ans['def'], ans['plan'] = tot('prod'), tot('good'), tot('defects'), tot('plan')
    ans['rate'] = ans['def'] / ans['prod']
    ans['ppm'] = ans['rate'] * 1e6
    ans['ach'] = ans['prod'] / ans['plan']
    rw = [r for r in rows if r['rework'] is not None]
    ans['fty'] = (sum(r['good'] - r['rework'] for r in rw)) / sum(r['prod'] for r in rw)
    ans['rw_blank'] = len(rows) - len(rw)
    run, load = tot('run'), len(rows) * LOAD_MIN
    perf = sum(r['ct'] * r['prod'] / 60 for r in rows) / run
    ans['oee'] = (run / load, perf, ans['good'] / ans['prod'], run / load * perf * ans['good'] / ans['prod'])
    ans['reason_blank'] = sum(1 for r in rows if r.get('reason_blank'))
    ans['lineshift'] = {}
    for li in LINES:
        for sh in SHIFTS:
            rs = [r for r in rows if r['line'] == li and r['shift'] == sh]
            ans['lineshift'][(li, sh)] = tot('defects', rs) / tot('prod', rs)
    ans['month'] = {m: tot('defects', [r for r in rows if r['d'].month == m]) / tot('prod', [r for r in rows if r['d'].month == m])
                    for m in range(1, 10)}
    by = {}
    for r in rows:
        if r['dtype']:
            by[r['dtype']] = by.get(r['dtype'], 0) + r['defects']
    total_def = sum(by.values())
    cum, pareto = 0, []
    for t, v in sorted(by.items(), key=lambda kv: -kv[1]):
        cum += v
        pareto.append((t, v, v / total_def, cum / total_def))
    ans['pareto'] = pareto
    in_win = lambda r: r['line'] == '2호라인' and r['shift'] == '야간' and date(2026, 6, 1) <= r['d'] <= date(2026, 7, 15)
    win = [r for r in rows if in_win(r)]
    oth = [r for r in rows if r['line'] == '2호라인' and r['shift'] == '야간' and not in_win(r)]
    dim = lambda rs: sum(r['defects'] for r in rs if r['dtype'] == '치수불량') / tot('prod', rs)
    ans['l2n_win'] = (len(win), dim(win), tot('defects', win) / tot('prod', win))
    ans['l2n_oth'] = (len(oth), dim(oth), tot('defects', oth) / tot('prod', oth))
    pl = [r for r in rows if r['fam'] == 'PL']
    bad = [r for r in pl if r['lot'] == BAD_LOT]
    ok = [r for r in pl if r['lot'] != BAD_LOT]
    bub = lambda rs: sum(r['defects'] for r in rs if r['dtype'] == '기포') / tot('prod', rs)
    ans['lot'] = (len(bad), bub(bad), tot('defects', bad) / tot('prod', bad), len(ok), bub(ok), tot('defects', ok) / tot('prod', ok))
    ans['bad_dates'] = sorted({r['d'] for r in bad})
    ans['bad_lines'] = sorted({(r['line'], r['shift']) for r in bad})
    daily = {}
    for r in rows:
        p, d = daily.get(r['d'], (0, 0))
        daily[r['d']] = (p + r['prod'], d + r['defects'])
    rates = {k: d / p for k, (p, d) in daily.items()}
    mu, sd = statistics.mean(rates.values()), statistics.stdev(rates.values())
    ans['daily'] = (mu, sd, sorted((k, v) for k, v in rates.items() if v > mu + 3 * sd))
    cell = {}
    for r in rows:
        if r['d'] == date(2026, 3, 18):
            cell[(r['line'], r['shift'])] = r['defects'] / r['prod']
    ans['mar18'] = cell[('3호라인', '주간')]
    ans['mar18_type'] = [r['dtype'] for r in rows if r['d'] == date(2026, 3, 18) and r['line'] == '3호라인' and r['shift'] == '주간'][0]
    ans['feb20'] = [r for r in rows if r['d'] == date(2026, 2, 20) and r['line'] == '2호라인' and r['shift'] == '주간'][0]
    return ans


# =============================================================== 2. 재고·출고
STOCK_ITEMS = [  # 코드, 품명, 단위, 일평균(영업일), 리드타임(일), 최소발주량, 발주단위, 단가(원), 목표판정, 특이사항
    ('R-001', 'ABS 수지(25kg포)', '포', 40, 14, 400, 40, 52000, 'late', ''),
    ('R-002', 'PP 수지(25kg포)', '포', 55, 14, 400, 40, 48000, 'now', ''),
    ('R-003', '알루미늄 잉곳', 'kg', 300, 21, 3000, 500, 3600, 'soon', ''),
    ('R-004', '구리 선재', 'kg', 120, 21, 1000, 100, 14500, 'late', ''),
    ('R-005', '포장 박스 A형', '장', 260, 7, 2000, 500, 420, 'ok', '재고0'),
    ('R-006', '포장 박스 B형', '장', 180, 7, 2000, 500, 380, 'now', '입고예정'),
    ('R-007', '완충 포장재(롤)', '롤', 12, 10, 60, 10, 18500, 'soon', ''),
    ('R-008', '방열 필름', '롤', 25, 14, 100, 10, 22000, 'ok', '여름성수기'),
    ('R-009', '단열 시트', '롤', 18, 14, 100, 10, 26000, 'now', '겨울성수기'),
    ('R-010', '볼트·너트 키트', '세트', 90, 10, 1000, 100, 1200, 'ok', ''),
    ('R-011', '고무 패킹', '개', 150, 10, 2000, 500, 380, 'soon', ''),
    ('R-012', '윤활유(18L)', '통', 6, 7, 20, 4, 68000, 'ok', '재고0'),
    ('R-013', '세척제(20L)', '통', 5, 7, 20, 4, 41000, 'ok', ''),
    ('R-014', '구형 세척제 A(20L)', '통', 0, 7, 20, 4, 39000, 'dormant', '90일 이상 출고 없음'),
    ('R-015', '도료(18L)', '통', 9, 14, 30, 6, 85000, 'ok', ''),
    ('R-016', '접착제(1kg)', '개', 30, 10, 120, 20, 9800, 'ok', '이력 1년 미만'),
    ('R-017', 'IoT 센서 모듈(신규)', '개', 20, 28, 200, 50, 14200, 'short', '신규 - 이력 짧음'),
    ('R-018', '구형 라벨 스티커', '권', 0, 7, 50, 10, 6500, 'dormant', '90일 이상 출고 없음'),
]
START, ASOF = date(2024, 7, 1), date(2026, 9, 30)
STOCKOUTS = {'R-005': (date(2025, 11, 10), date(2025, 11, 14)), 'R-012': (date(2026, 9, 15), date(2026, 9, 18))}
FIRST_USE = {'R-016': date(2025, 11, 3), 'R-017': date(2026, 9, 1)}
LAST_USE = {'R-014': date(2026, 5, 28), 'R-018': date(2026, 6, 19)}
SEASON = {'R-008': (0.8, date(2026, 7, 15)), 'R-009': (0.8, date(2026, 1, 10))}
WDAY = {0: 1.15, 1: 1.0, 2: 1.0, 3: 0.95, 4: 0.9}
Z = 1.65


def make_usage(rng):
    days = [d for d in daterange(START, ASOF) if is_workday(d)]
    series = {}
    for code, _n, _u, base, *_rest in STOCK_ITEMS:
        s = {}
        for i, d in enumerate(days):
            if code in FIRST_USE and d < FIRST_USE[code]:
                s[d] = None
                continue
            if code in LAST_USE and d > LAST_USE[code]:
                s[d] = 0
                continue
            lvl = base * (1 + 0.0002 * i) * WDAY[d.weekday()]
            if code in SEASON:
                amp, peak = SEASON[code]
                lvl *= 1 + amp * math.cos(2 * math.pi * ((d - peak).days % 365) / 365)
            if code == 'R-017':
                lvl *= 1 + 0.01 * (d - FIRST_USE[code]).days
            q = max(0, round(rng.gauss(lvl, lvl * 0.28)))
            if code in STOCKOUTS and STOCKOUTS[code][0] <= d <= STOCKOUTS[code][1]:
                q = 0
            if code in LAST_USE and d == LAST_USE[code]:
                q = max(q, 1)
            s[d] = q
        series[code] = s
    return days, series


def stockout_days(code):
    if code not in STOCKOUTS:
        return set()
    return set(daterange(*STOCKOUTS[code]))


def avg_daily(series, code, a, b):
    """달력일 기준 일평균 = 기간 출고합 / (기간 달력일수 - 재고0 일수)."""
    so = stockout_days(code)
    days = [d for d in daterange(a, b) if d not in so]
    if not days:
        return 0.0
    return sum(series[code].get(d) or 0 for d in days) / len(days)


def month_workday_avg(series, code, y, m):
    so = stockout_days(code)
    vals = [v for d, v in series[code].items() if d.year == y and d.month == m and v is not None and d not in so]
    return statistics.mean(vals) if vals else 0.0


def analyze_item(series, item, asof=ASOF):
    code, name, unit, _b, lead, moq, step, price, _t, _note = item
    s = series[code]
    first = next((d for d, v in s.items() if v is not None), None)
    hist = (asof - first).days + 1
    out = dict(code=code, name=name, hist=hist, lead=lead, moq=moq, step=step, price=price)
    last_use = max((d for d, v in s.items() if v), default=None)
    out['days_since'] = (asof - last_use).days if last_use else None
    if hist < 56:
        out['status'] = '이력 부족'
        return out
    if out['days_since'] is not None and out['days_since'] >= 90:
        out['status'] = '장기 체화'
        return out
    d4 = avg_daily(series, code, asof - timedelta(days=27), asof)
    d1 = avg_daily(series, code, asof - timedelta(days=6), asof)
    so = stockout_days(code)
    weeks = []
    for k in range(12):
        b = asof - timedelta(days=7 * k)
        a = b - timedelta(days=6)
        if any(d in so for d in daterange(a, b)):
            continue
        weeks.append(sum(s.get(d) or 0 for d in daterange(a, b)))
    sigma_w = statistics.stdev(weeks)
    factor = 1.0
    out['season'] = False
    if hist >= 365:
        tm = (asof + timedelta(days=lead))
        cur = month_workday_avg(series, code, asof.year - 1, asof.month)
        nxt = month_workday_avg(series, code, tm.year - 1, tm.month)
        if cur > 0 and nxt > 0:
            factor, out['season'] = nxt / cur, True
    d = d4 * factor
    ss = Z * sigma_w * math.sqrt(lead / 7)
    rop = d * lead + ss
    out.update(d4=d4, d1=d1, sigma_w=sigma_w, factor=factor, d=d, ss=ss, rop=rop)
    return out


def decide(o, avail):
    days = avail / o['d']
    if days < o['lead']:
        st = '이미 늦음'
    elif avail <= o['rop']:
        st = '지금 발주'
    elif (avail - o['rop']) / o['d'] <= 14:
        st = '곧 발주'
    else:
        st = '정상'
    target = o['rop'] + o['d'] * 30
    qty = max(o['moq'], target - avail)
    qty = math.ceil(qty / o['step']) * o['step']
    return st, days, qty


def set_stock(item, o, rng):
    """목표 판정이 나오도록 현재고(와 입고예정)를 정한다."""
    target = item[8]
    d, L, ss, rop = o['d'], o['lead'], o['ss'], o['rop']
    avail = {'late': 0.6 * d * L, 'now': d * L + 0.5 * ss, 'soon': rop + 7 * d, 'ok': rop + 50 * d}[target]
    pending = 0
    if item[9] == '입고예정':
        pending = int(rop + 30 * d)
        stock = max(0, int(avail * 0.2))
        # 현재고만 보면 늦음, 입고예정을 더하면 정상이 되도록
        avail = stock + pending
        return stock, pending, avail
    stock = int(avail)
    if target == 'soon' and rng.random() < 0.5:
        pending = int(stock * 0.3)
        stock -= pending
    return stock, pending, stock + pending


def write_stock(days, series, infos):
    wb = Workbook()
    ws = wb.active
    ws.title = '일별출고'
    ws.append(['일자'] + [i[0] for i in STOCK_ITEMS])
    for d in days:
        ws.append([d] + [series[i[0]][d] for i in STOCK_ITEMS])
        ws.cell(row=ws.max_row, column=1).number_format = 'yyyy-mm-dd'
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = 'B2'
    ws1 = wb.create_sheet('품목정보')
    ws1.append(['품목코드', '품명', '단위', '리드타임(일)', '최소발주량', '발주단위', '단가(원)'])
    for i in STOCK_ITEMS:
        ws1.append([i[0], i[1], i[2], i[4], i[5], i[6], i[7]])
    ws2 = wb.create_sheet('현재고')
    ws2.append(['품목코드', '품명', '기준일', '현재고'])
    for i in STOCK_ITEMS:
        ws2.append([i[0], i[1], ASOF, infos[i[0]]['stock']])
        ws2.cell(row=ws2.max_row, column=3).number_format = 'yyyy-mm-dd'
    ws3 = wb.create_sheet('입고예정')
    ws3.append(['품목코드', '품명', '입고예정일', '수량'])
    for i in STOCK_ITEMS:
        p = infos[i[0]]['pending']
        if p:
            ws3.append([i[0], i[1], ASOF + timedelta(days=min(7, i[4] // 2)), p])
            ws3.cell(row=ws3.max_row, column=3).number_format = 'yyyy-mm-dd'
    ws4 = wb.create_sheet('재고0기간')
    ws4.append(['품목코드', '시작일', '종료일', '비고'])
    for c, (a, b) in STOCKOUTS.items():
        ws4.append([c, a, b, '재고가 없어 출고하지 못한 기간(출고 0은 수요 0이 아님)'])
        for col in (2, 3):
            ws4.cell(row=ws4.max_row, column=col).number_format = 'yyyy-mm-dd'
    ws5 = wb.create_sheet('안내')
    for t in ['가상 데이터입니다. 공휴일에는 출고가 없어 행 자체가 없습니다(주말도 없음).',
              '빈칸은 아직 취급하기 전(신규 품목)이라는 뜻이며 0과 다릅니다.',
              f'현재고·입고예정 기준일은 {ASOF}.']:
        ws5.append([t])
    wb.save(OUT / '재고_출고이력.xlsx')


# =============================================================== README
def pct(x, n=2):
    return f'{x * 100:.{n}f}%'


def write_readme(pa, results, stock_days):
    L = []
    A = L.append
    A('# 제조 실습용 가상 데이터')
    A('')
    A('모든 회사·품목·숫자는 **가상**입니다(`python3 samples/make_samples.py`로 다시 만들 수 있고 값은 항상 같습니다). 정답 수치는 스크립트가 계산해서 적었습니다.')
    A('공휴일(관공서의 공휴일에 관한 규정에 따른 공휴일 표)에는 생산·출고가 없습니다.')
    A('')
    A('| 파일 | 쓰는 스킬 |')
    A('|---|---|')
    A('| 생산실적.xlsx | production-defect-analysis |')
    A('| 재고_출고이력.xlsx | reorder-timing |')
    A('')
    A('## 1. 생산실적.xlsx (production-defect-analysis)')
    A('')
    A(f'- {pa["n"]}행. 2026-01~09의 월·수·금(공휴일 제외)만 기록(무료 플랜 행 수 한도 고려). 라인 3개 x 주·야 2교대, 품목 5개(M-101~M-105), 자재 로트 열 있음. 교대당 부하시간 {LOAD_MIN}분(가동시간 + 비가동시간).')
    A('- 불량유형은 줄마다 "주된 유형" 하나만 적혀 있어 파레토는 이 기준입니다. 정의: 생산수량 = 양품수량 + 불량수량, 양품수량은 재작업 후 합격분 포함, 재작업수량은 그중 재작업을 거친 수.')
    A('- **일부러 넣은 지저분함**: 일자 5줄은 문자열("2026.03.04"), 라인 표기 8줄이 "1라인"처럼 다름, 비가동사유 빈칸 %d줄(비가동시간은 있음), 재작업수량 빈칸 %d줄(0이 아님).' % (pa['reason_blank'], pa['rw_blank']))
    A('- **심어 둔 패턴 3가지**')
    A('  1. 2호라인 **야간조**, 2026-06-01~07-15에 **치수불량**이 늚 (설비 쪽 가설 - 원인은 데이터로 확정 불가, 확인 필요).')
    A(f'  2. 자재 로트 **{BAD_LOT}**(수지 소재, 5월 하순)을 쓴 줄(M-102·M-103, 라인·교대 무관)에서 **기포**가 늚 - 라인·교대가 아니라 로트가 공통 원인.')
    A('  3. 2026-03-18 3호라인 주간 하루에만 **이물** 불량 급증(하루짜리 이상치). 2026-02-20 2호라인 주간은 설비고장으로 비가동 240분(불량 문제는 아님).')
    A('')
    A('**검산용 정답** (전체 행 기준, 빈칸 줄은 해당 지표 계산에서 제외)')
    A('')
    A(f'- 생산 {pa["prod"]:,} / 양품 {pa["good"]:,} / 불량 {pa["def"]:,} / 계획 {pa["plan"]:,}개')
    A(f'- 불량률 {pct(pa["rate"])} ({pa["ppm"]:,.0f}ppm), 계획 달성률 {pct(pa["ach"], 1)}')
    A(f'- 직행률 {pct(pa["fty"])} (= (양품 - 재작업) / 생산, 재작업 빈칸 {pa["rw_blank"]}줄 제외. 분모를 투입·완료 중 무엇으로 볼지 회사 정의 확인 필요)')
    o = pa['oee']
    A(f'- OEE {pct(o[3], 1)} = 시간가동률 {pct(o[0], 1)} x 성능가동률 {pct(o[1], 1)} x 양품률 {pct(o[2], 1)} (부하시간 {LOAD_MIN}분, 기준 사이클타임은 품목기준 시트, 가동시간 = 부하 - 비가동. 회사 정의가 다르면 그것을 사용)')
    A('')
    A('| 불량률 | 주간 | 야간 |')
    A('|---|---|---|')
    for li in LINES:
        A(f"| {li} | {pct(pa['lineshift'][(li, '주간')])} | {pct(pa['lineshift'][(li, '야간')])} |")
    A('')
    A('월별 불량률: ' + ', '.join(f'{m}월 {pct(v)}' for m, v in pa['month'].items()))
    A('')
    A('불량유형 파레토 (불량수량 기준)')
    A('')
    A('| 불량유형 | 불량수량 | 비중 | 누적 |')
    A('|---|---|---|---|')
    for t, v, p, c in pa['pareto']:
        A(f'| {t} | {v:,} | {pct(p, 1)} | {pct(c, 1)} |')
    A('')
    w, oth = pa['l2n_win'], pa['l2n_oth']
    A(f'- 2호라인 야간: 6/1~7/15({w[0]}줄) 치수불량률 {pct(w[1])}, 전체 불량률 {pct(w[2])} / 그 외 기간({oth[0]}줄) 치수불량률 {pct(oth[1])}, 전체 {pct(oth[2])}')
    b = pa['lot']
    A(f'- 수지 소재 품목: {BAD_LOT} 사용 {b[0]}줄 기포율 {pct(b[1])}(전체 불량률 {pct(b[2])}) / 다른 로트 {b[3]}줄 기포율 {pct(b[4])}(전체 {pct(b[5])}). {BAD_LOT} 사용일: {pa["bad_dates"][0]}~{pa["bad_dates"][-1]}, 걸친 라인·교대 {len(pa["bad_lines"])}곳')
    mu, sd, outl = pa['daily']
    A(f'- 일별 전체 불량률 평균 {pct(mu)}, 표준편차 {pct(sd)} -> 평균+3표준편차({pct(mu + 3 * sd)}) 초과일: ' + (', '.join(f'{k}({pct(v)})' for k, v in outl) or '없음') + ' (3/18은 심어 둔 하루짜리 이상치, 5/18·5/20은 불량 로트 PL-2605-B 사용일)')
    A(f'- 2026-03-18 3호라인 주간 불량률 {pct(pa["mar18"], 1)}, 주된 유형 {pa["mar18_type"]}')
    f = pa['feb20']
    A(f'- 2026-02-20 2호라인 주간: 비가동 {f["down"]}분({f["reason"]}), 계획 달성률 {pct(f["prod"] / f["plan"], 1)}')
    A('')
    A('## 2. 재고_출고이력.xlsx (reorder-timing)')
    A('')
    A(f'- `일별출고` 시트 {len(stock_days)}행(품목 18열, 2024-07-01~2026-09-30, 영업일만), `품목정보`(리드타임·최소발주량·발주단위·단가), `현재고`, `입고예정`, `재고0기간` 시트. 기준일은 2026-09-30.')
    A('- **심어 둔 함정**: ① 계절성 품목 R-008(여름 성수기)·R-009(겨울 성수기) ② 재고 0 기간(R-005 2025-11-10~14, R-012 2026-09-15~18 - 이 날의 출고 0은 수요 0이 아니므로 평균에서 제외, 최근 4주 안에 걸린 R-012가 핵심) ③ 이력이 짧은 신규 R-017(2026-09-01 시작, 그 전은 빈칸), 이력 1년 미만 R-016 ④ 90일 이상 출고 없는 R-014·R-018 ⑤ 1주 평균은 추석 연휴 주(9/24·25 휴무)라 낮게 나옴 ⑥ R-006은 현재고만 보면 부족하지만 입고예정을 더하면 충분함.')
    A('')
    A('**검산용 정답 - 스킬 본문의 공식 그대로** (기준일 2026-09-30)')
    A('- 일평균(4주) = 최근 28일 출고합 / (28일 - 재고0 일수) (달력일 기준, 휴무일은 출고 0으로 포함). 1주는 최근 7일.')
    A('- 주간 변동 = 최근 12주 각 주(7일) 출고량의 표준편차(재고0이 낀 주 제외). 안전재고 = 1.65 x 주간 변동 x √(리드타임/7) (서비스수준 95%, 정규분포 가정).')
    A('- 계절성 보정(이력 1년 이상일 때만) = 전년 (기준일+리드타임)이 속한 달의 영업일 평균 / 전년 기준일이 속한 달의 영업일 평균. 보정 일평균 = 4주 일평균 x 보정.')
    A('- 발주점 = 보정 일평균 x 리드타임 + 안전재고, 가용재고 = 현재고 + 입고예정, 재고일수 = 가용재고 / 보정 일평균.')
    A('- 판정: 재고일수 < 리드타임 -> 이미 늦음 / 가용재고 <= 발주점 -> 지금 발주 / (가용 - 발주점) / 일평균 <= 14일 -> 곧 발주 / 그 외 정상. 발주량 = max(최소발주량, 발주점 + 30일치 - 가용재고)를 발주단위로 올림.')
    A('')
    A('| 품목 | 일평균(4주) | 일평균(1주) | 보정 | 안전재고 | 발주점 | 현재고 | 입고예정 | 재고일수 | 리드타임 | 판정 | 발주량 |')
    A('|---|---|---|---|---|---|---|---|---|---|---|---|')
    cnt = {}
    held = []
    for r in results:
        if 'd' not in r:
            A(f"| {r['code']} {r['name']} | - | - | - | - | - | {r['stock']:,} | {r['pending']:,} | - | {r['lead']} | {r['status']} | - |")
            if r['status'] == '장기 체화':
                held.append(r)
            continue
        cnt[r['verdict']] = cnt.get(r['verdict'], 0) + 1
        fac = f"x{r['factor']:.2f}" if r['season'] else '없음'
        A(f"| {r['code']} {r['name']} | {r['d4']:.1f} | {r['d1']:.1f} | {fac} | {r['ss']:.0f} | {r['rop']:.0f} | {r['stock']:,} | {r['pending']:,} | {r['days']:.1f}일 | {r['lead']}일 | {r['verdict']} | {r['qty']:,} |")
    A('')
    A('판정 건수: ' + ', '.join(f'{k} {v}' for k, v in cnt.items()) + ', 이력 부족 1, 장기 체화 2')
    A('')
    A('장기 체화(90일 이상 출고 없음):')
    for r in held:
        A(f"- {r['code']} {r['name']}: 마지막 출고 후 {r['days_since']}일, 현재고 {r['stock']:,} x 단가 {r['price']:,}원 = {r['stock'] * r['price']:,}원 묶임")
    A(f"- 합계 {sum(r['stock'] * r['price'] for r in held):,}원")
    A('')
    A('## 주의')
    A('연습용 파일입니다. **실제 회사 데이터를 무료·개인용 AI 계정에 올리지 마세요.**')
    (OUT / 'README.md').write_text('\n'.join(L) + '\n', encoding='utf-8')


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rng = random.Random(SEED)
    rows = make_production(rng)
    write_production(rows)
    pa = analyze_production(rows)

    days, series = make_usage(rng)
    infos, results = {}, []
    for item in STOCK_ITEMS:
        o = analyze_item(series, item)
        if 'd' in o:
            stock, pending, avail = set_stock(item, o, rng)
            o['verdict'], o['days'], o['qty'] = decide(o, avail)
        else:
            stock = 600 if item[0] == 'R-014' else 900 if item[0] == 'R-018' else 300
            if item[0] == 'R-017':
                stock = 700
            pending = 0
        o['stock'], o['pending'] = stock, pending
        infos[item[0]] = dict(stock=stock, pending=pending)
        results.append(o)
    write_stock(days, series, infos)
    write_readme(pa, results, days)
    print('생성 완료:', OUT)


if __name__ == '__main__':
    main()
